from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from flask import abort, redirect, render_template, request, session, url_for
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
ALLOWED_EXTENSIONS = {".xlsx"}

REQUIRED_COLUMNS = {
    "Site ID",
    "Site Name",
    "Long",
    "Lat",
    "City",
    "Kecamatan",
    "Program Name",
    "RAN VENDOR",
    "Update status",
}

def _txt(v):
    return "" if v is None else str(v).strip()

def _date(v):
    if v in (None, ""):
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date().isoformat()
        except Exception:
            return _txt(v)
    return _txt(v)

def _load_json(path: Path):
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)

def _atomic_write(path: Path, data: dict):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)

def _summary(data: dict):
    s = data.get("summary", {})
    st = s.get("status_counts", {})
    return {
        "total_sites": int(s.get("total_sites", 0) or 0),
        "on_air": int(st.get("On Air", 0) or 0),
        "integrated": int(st.get("Integrated - NY OA", 0) or 0),
        "on_progress": int((st.get("Installation", 0) or 0) + (st.get("NY Install", 0) or 0)),
    }
def convert_xlsx(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Sitelist" not in wb.sheetnames:
            raise ValueError("Sheet 'Sitelist' tidak ditemukan.")

        ws = wb["Sitelist"]
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            raise ValueError("Sheet Sitelist kosong.")

        idx = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}
        missing = sorted(REQUIRED_COLUMNS - set(idx))
        if missing:
            raise ValueError("Kolom wajib tidak ditemukan: " + ", ".join(missing))

        def cell(row, *names, default=""):
            for name in names:
                if name in idx:
                    v = row[idx[name]]
                    if v not in (None, ""):
                        return v
            return default

        sites = []

        for row in rows:
            site_id = _txt(cell(row, "Site ID"))
            if not site_id:
                continue

            update_status = _txt(cell(row, "Update status"))
            status = "NY Install" if update_status == "On Progress" else update_status

            try:
                lng = float(cell(row, "Long", "Lng", "Longitude"))
            except Exception:
                lng = None

            try:
                lat = float(cell(row, "Lat", "Latitude"))
            except Exception:
                lat = None

            sites.append({
                "no": cell(row, "No", default=""),
                "site_id": site_id,
                "site_name": _txt(cell(row, "Site Name")),
                "lng": lng,
                "lat": lat,
                "city": _txt(cell(row, "City")),
                "kecamatan": _txt(cell(row, "Kecamatan")),
                "desa": _txt(cell(row, "Desa")),
                "program": _txt(cell(row, "Program Name", "Program")),
                "on_air_date": _date(cell(row, "On Air Date")),
                "integration_date": _date(cell(row, "Integration Date", "Integrated Date")),
                "vendor": _txt(cell(row, "RAN VENDOR", "Vendor")),
                "sub_region": _txt(cell(row, "Sub Region")),
                "lonlat": _txt(cell(row, "Lonlat")),
                "update_status": update_status,
                "status": status,
                "week_to_closed": _txt(cell(row, "Week to Closed")),
            })
    finally:
        wb.close()

    if not sites:
        raise ValueError("Tidak ada Site ID valid di sheet Sitelist.")

    status_counts = Counter(s["status"] for s in sites)
    for key in ("On Air", "Integrated - NY OA", "Installation", "NY Install"):
        status_counts.setdefault(key, 0)

    update_counts = Counter(s["update_status"] for s in sites)
    vendor_counts = Counter(s["vendor"] for s in sites if s["vendor"])
    program_counts = Counter(s["program"] for s in sites if s["program"])
    subregion_counts = Counter(s["sub_region"] for s in sites if s["sub_region"])
    week_counts = Counter(s["week_to_closed"] for s in sites if s["week_to_closed"])
    city_counts = Counter(s["city"] for s in sites if s["city"])

    groups = defaultdict(list)
    for s in sites:
        groups[s["city"]].append(s)

    cities = []
    for city, group in sorted(groups.items(), key=lambda kv: (-len(kv[1]), kv[0] or "")):
        sc = Counter(s["status"] for s in group)
        vc = Counter(s["vendor"] for s in group)
        total = len(group)
        on_air = sc.get("On Air", 0)
        kec_count = len({s["kecamatan"] for s in group if s["kecamatan"]})
        desa_count = len({s["desa"] for s in group if s["desa"]})

        cities.append({
            "city": city,
            "total": total,
            "on_air": on_air,
            "integrated_ny_oa": sc.get("Integrated - NY OA", 0),
            "installation": sc.get("Installation", 0),
            "ny_install": sc.get("NY Install", 0),
            "on_progress": sum(1 for s in group if s["update_status"] == "On Progress"),
            "on_air_rate": (on_air / total * 100) if total else 0,
            "nokia": vc.get("Nokia", 0),
            "huawei": vc.get("Huawei", 0),
            "kecamatan": kec_count,
            "kecamatan_count": kec_count,
            "desa_count": desa_count,
        })

    site_id_counts = Counter(s["site_id"] for s in sites)

    duplicate_site_ids = sorted(
        [
            {
                "site_id": site_id,
                "count": count,
                "rows": [
                    {
                        "site_name": s.get("site_name", ""),
                        "city": s.get("city", ""),
                        "kecamatan": s.get("kecamatan", ""),
                        "status": s.get("status", ""),
                        "vendor": s.get("vendor", ""),
                        "program": s.get("program", ""),
                        "lat": s.get("lat"),
                        "lng": s.get("lng"),
                    }
                    for s in sites
                    if s.get("site_id") == site_id
                ],
            }
            for site_id, count in site_id_counts.items()
            if count > 1
        ],
        key=lambda x: (-x["count"], x["site_id"])
    )

    dup_count = len(duplicate_site_ids)

    # Exact duplicates: same meaningful data, ignoring Excel row number / No.
    def duplicate_fingerprint(site):
        return (
            site.get("site_id", ""),
            site.get("site_name", ""),
            site.get("lng"),
            site.get("lat"),
            site.get("city", ""),
            site.get("kecamatan", ""),
            site.get("desa", ""),
            site.get("program", ""),
            site.get("on_air_date", ""),
            site.get("integration_date", ""),
            site.get("vendor", ""),
            site.get("sub_region", ""),
            site.get("lonlat", ""),
            site.get("update_status", ""),
            site.get("status", ""),
            site.get("week_to_closed", ""),
        )

    fingerprint_counts = Counter(duplicate_fingerprint(s) for s in sites)
    exact_duplicate_groups = []
    exact_duplicate_row_count = 0

    for fingerprint, count in fingerprint_counts.items():
        if count <= 1:
            continue

        matching = [
            s for s in sites
            if duplicate_fingerprint(s) == fingerprint
        ]
        first = matching[0]
        exact_duplicate_row_count += count - 1

        exact_duplicate_groups.append({
            "site_id": first.get("site_id", ""),
            "site_name": first.get("site_name", ""),
            "city": first.get("city", ""),
            "kecamatan": first.get("kecamatan", ""),
            "status": first.get("status", ""),
            "vendor": first.get("vendor", ""),
            "program": first.get("program", ""),
            "count": count,
        })

    exact_duplicate_groups.sort(
        key=lambda x: (-x["count"], x["site_id"], x["site_name"])
    )

    invalid_coords = sum(
        1 for s in sites
        if s["lat"] is None
        or s["lng"] is None
        or not (-90 <= s["lat"] <= 90)
        or not (-180 <= s["lng"] <= 180)
    )

    return {
        "meta": {
            "source_file": path.name,
            "source_date": datetime.now().date().isoformat(),
            "generated_for": "East Java Vanguard",
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "note": (
                "Generated by Data Update Center. "
                "On Progress mapped to NY Install for dashboard compatibility."
            ),
        },
        "summary": {
            "total_sites": len(sites),
            "status_counts": dict(status_counts),
            "update_status_counts": dict(update_counts),
            "vendor_counts": dict(vendor_counts),
            "program_counts": dict(program_counts),
            "sub_region_counts": dict(subregion_counts),
            "week_to_closed_counts": dict(week_counts),
            "city_counts": dict(city_counts),
        },
        "cities": cities,
        "sites": sites,
        "_validation": {
            "duplicate_site_id_count": dup_count,
            "duplicate_site_ids": duplicate_site_ids,
            "exact_duplicate_group_count": len(exact_duplicate_groups),
            "exact_duplicate_row_count": exact_duplicate_row_count,
            "exact_duplicate_groups": exact_duplicate_groups,
            "invalid_coordinate_count": invalid_coords,
            "apply_blocked": bool(exact_duplicate_groups),
        },
    }


def register_update_center(app, require_role, verify_csrf, csrf_token, audit):
    incoming = Path("data/incoming")
    staging = Path("data/staging")
    backups = Path("data/backups")

    incoming.mkdir(parents=True, exist_ok=True)
    staging.mkdir(parents=True, exist_ok=True)
    backups.mkdir(parents=True, exist_ok=True)

    data_path = Path("data/Vanguard.json")
    if not data_path.exists() and Path("data/vanguard.json").exists():
        data_path = Path("data/vanguard.json")

    app.config.setdefault("MAX_CONTENT_LENGTH", MAX_UPLOAD_BYTES)

    @app.route("/admin/data-update", methods=["GET"])
    @require_role("admin")
    def data_update_center():
        files = sorted(
            [p for p in incoming.iterdir() if p.is_file()],
            key=lambda p: p.stat().st_mtime,
            reverse=True
        )
        pending = []
        for p in files[:50]:
            pending.append({
                "name": p.name,
                "size_mb": round(p.stat().st_size / 1024 / 1024, 2),
                "uploaded_at": datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
            })

        audit("DATA_UPDATE_CENTER_VIEW")
        return render_template(
            "data_update_center.html",
            user=session["user"],
            page="data_update",
            csrf=csrf_token(),
            pending=pending,
        )

    @app.route("/admin/data-update/upload", methods=["POST"])
    @require_role("admin")
    def data_update_upload():
        verify_csrf()
        f = request.files.get("dataset")
        if not f or not f.filename:
            abort(400, "File belum dipilih.")

        filename = secure_filename(f.filename)
        if Path(filename).suffix.lower() not in ALLOWED_EXTENSIONS:
            abort(400, "Hanya file .xlsx yang diizinkan.")

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target = incoming / f"{stamp}_{filename}"
        f.save(target)

        if target.stat().st_size > MAX_UPLOAD_BYTES:
            target.unlink(missing_ok=True)
            abort(413, "Ukuran file maksimal 20 MB.")

        audit("DATA_UPDATE_FILE_STAGED", target.name)
        return redirect(url_for("data_update_center", uploaded="1"))

    @app.route("/admin/data-update/validate", methods=["POST"])
    @require_role("admin")
    def data_update_validate():
        verify_csrf()
        name = request.form.get("name", "").strip()

        if not name or Path(name).name != name:
            abort(400)

        source = incoming / name
        if not source.exists():
            abort(404)

        try:
            new_data = convert_xlsx(source)
        except Exception as e:
            audit("DATA_UPDATE_VALIDATE_FAILED", f"{name}:{e}")
            return render_template(
                "data_update_error.html",
                user=session["user"],
                page="data_update",
                message=str(e),
            ), 400

        token = uuid.uuid4().hex
        stage_json = staging / f"{token}.json"
        stage_meta = staging / f"{token}.meta.json"

        current = _load_json(data_path)
        before = _summary(current)
        after = _summary(new_data)

        meta = {
            "token": token,
            "source_name": name,
            "before": before,
            "after": after,
            "delta": {k: after[k] - before.get(k, 0) for k in after},
            "validation": new_data.get("_validation", {}),
        }

        _atomic_write(stage_json, new_data)
        _atomic_write(stage_meta, meta)

        audit("DATA_UPDATE_VALIDATED", f"{name}:{after['total_sites']}")
        return redirect(url_for("data_update_preview", token=token))

    @app.route("/admin/data-update/preview/<token>", methods=["GET"])
    @require_role("admin")
    def data_update_preview(token):
        if not re.fullmatch(r"[a-f0-9]{32}", token):
            abort(400)

        stage_meta = staging / f"{token}.meta.json"
        if not stage_meta.exists():
            abort(404)

        preview = _load_json(stage_meta)

        return render_template(
            "data_update_preview.html",
            user=session["user"],
            page="data_update",
            preview=preview,
            csrf=csrf_token(),
        )

    @app.route("/admin/data-update/apply", methods=["POST"])
    @require_role("admin")
    def data_update_apply():
        verify_csrf()
        token = request.form.get("token", "").strip()

        if not re.fullmatch(r"[a-f0-9]{32}", token):
            abort(400)

        stage_json = staging / f"{token}.json"
        stage_meta = staging / f"{token}.meta.json"

        if not stage_json.exists() or not stage_meta.exists():
            abort(404)

        new_data = _load_json(stage_json)
        meta = _load_json(stage_meta)

        validation = meta.get("validation", {})
        if validation.get("apply_blocked"):
            audit(
                "DATASET_UPDATE_BLOCKED_EXACT_DUPLICATES",
                (
                    f"{meta.get('source_name', '')}:"
                    f"{validation.get('exact_duplicate_group_count', 0)} groups"
                ),
            )
            abort(
                409,
                "Apply diblokir karena ditemukan exact duplicate rows. "
                "Perbaiki dataset lalu Validate kembali."
            )

        new_data.pop("_validation", None)

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        if data_path.exists():
            shutil.copy2(data_path, backups / f"Vanguard_{stamp}.json")

        _atomic_write(data_path, new_data)

        source_name = meta.get("source_name")
        if source_name:
            try:
                (incoming / source_name).unlink(missing_ok=True)
            except PermissionError as exc:
                if getattr(exc, "winerror", None) != 32:
                    raise
                # The dataset is committed; a locked upload must not undo success.
                app.logger.warning(
                    "Dataset applied, but uploaded file is still locked: %s",
                    source_name,
                )
                audit("DATA_UPDATE_CLEANUP_DEFERRED", source_name)

        stage_json.unlink(missing_ok=True)
        stage_meta.unlink(missing_ok=True)

        audit(
            "DATASET_UPDATE_APPLIED",
            f"{source_name}:{meta.get('before',{}).get('total_sites',0)}->{meta.get('after',{}).get('total_sites',0)}"
        )

        return redirect(url_for("data_update_center", updated="1"))

    @app.route("/admin/data-update/delete", methods=["POST"])
    @require_role("admin")
    def data_update_delete():
        verify_csrf()
        name = request.form.get("name", "").strip()

        if not name or Path(name).name != name:
            abort(400)

        target = incoming / name
        if target.exists():
            target.unlink()
            audit("DATA_UPDATE_FILE_REMOVED", name)

        return redirect(url_for("data_update_center"))
